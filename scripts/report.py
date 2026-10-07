"""彙整比對結果 → Excel（總表＋變更明細）。
用法: python report.py --work <工作資料夾> --xlsx <輸出 Excel 路徑>
"""
import argparse, glob, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import cadcom

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--work", required=True); ap.add_argument("--xlsx", required=True)
    a = ap.parse_args()
    old_root, new_root = open(os.path.join(a.work, "roots.txt"), encoding="utf-8").read().split("\n")[:2]
    rel = lambda p, root: os.path.relpath(p, root).replace("\\", "/") if p else ""
    changed_pairs = {os.path.normcase(os.path.normpath(n)) for _, _, n in cadcom.read_lines(os.path.join(a.work, "pairs.txt"))}
    marked = {}
    for f in glob.glob(os.path.join(a.work, "res", "*_marked.json")):
        j = json.load(open(f, encoding="utf-8")); marked[os.path.normcase(os.path.normpath(j["new"]))] = j

    wb = Workbook(); ws = wb.active; ws.title = "總表"
    ws.append(["項次", "專業", "圖檔（新版，相對路徑）", "比對結果", "標示處數", "新增配置頁", "刪除配置頁", "標示檔", "備註"])
    dt = wb.create_sheet("變更明細")
    dt.append(["圖檔", "空間/配置", "標示編號", "變更說明", "新增物件數", "刪除物件數", "舊版文字", "新版文字", "相關圖塊", "座標範圍(x1,y1,x2,y2)"])
    n = 0
    for k, o, nf in cadcom.read_lines(os.path.join(a.work, "pairs_all.txt")):
        n += 1; r = rel(nf, new_root)
        if not os.path.exists(o):
            ws.append([n, k, r, "新增圖說（舊版無此圖）"]); continue
        m = marked.get(os.path.normcase(os.path.normpath(nf)))
        if m is None:
            pending = os.path.normcase(os.path.normpath(nf)) in changed_pairs   # make_pairs 已判定內容不同
            ws.append([n, k, r, "未完成比對（轉檔或比對失敗，需人工比對）" if pending else "相同"]); continue
        changed = m["rows"] or m["layouts_added"] or m["layouts_removed"]
        ws.append([n, k, r, "有變更" if changed else "內容無差異（僅存檔資訊不同）",
                   sum(1 for x in m["rows"] if x["no"] != "-"), "、".join(m["layouts_added"]), "、".join(m["layouts_removed"]),
                   rel(m["marked"], new_root), "；".join(m.get("notes", []))])
        for x in m["rows"]:
            dt.append([os.path.basename(nf), x["space"], x["no"], x["label"], x["n_add"], x["n_del"],
                       "\n".join(x["del_text"][:30]), "\n".join(x["add_text"][:30]), "、".join(x["blocks"][:10]), str(x["box"])])
    rm = os.path.join(a.work, "removed.txt")
    for k, o in (cadcom.read_lines(rm) if os.path.exists(rm) else []):
        n += 1; ws.append([n, k, rel(o, old_root), "新版資料夾沒有此圖（可能刪除或本次未發圖；路徑為舊版）"])

    for sh, widths in ((ws, [6, 10, 70, 30, 9, 20, 20, 60, 40]), (dt, [34, 16, 9, 44, 10, 10, 40, 40, 30, 40])):
        for i, w in enumerate(widths, 1): sh.column_dimensions[get_column_letter(i)].width = w
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for row in sh.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
        sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
    os.makedirs(os.path.dirname(os.path.abspath(a.xlsx)), exist_ok=True)
    wb.save(a.xlsx)
    print("Excel：", a.xlsx)

if __name__ == "__main__":
    main()
